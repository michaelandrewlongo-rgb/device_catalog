"""Export a color-coded Excel workbook: one sheet per device category.

Built from the same validated rows as the flat CSV and nested tree. Each sheet
shows only the specification columns that category actually uses, colored by
factor group so ID / OD / lengths / implant sizes / compatibility can be read
at a glance:

- blue   inner diameter (ID)
- green  outer diameters (OD, distal/proximal/middle, generic diameter)
- orange lengths (usable/stent/balloon)
- purple implant and device sizes (coil, stent, balloon, implant, wire, vessel)
- yellow compatibility (guidewire, catheter)
- grey   provenance (evidence layer, year, locator)

Values are the printed source values (decimal-normalized), same as the CSV.

    python -m pipeline.export_devices_workbook [--out PATH]
"""

from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .export_flat_dimensions import build_rows
from .knowledge_v2.validate_dimensions import summarize, validate_csv_rows

DEFAULT_OUT = Path.home() / "Downloads" / "device_dimensions_by_category.xlsx"

IDENTITY_COLUMNS = ["device_name", "manufacturer", "catalog_number", "product_name"]

# factor group -> (columns, fill color)
GROUPS = {
    "id": (["inner_diameter"], "DCE6F1"),                                   # blue
    "od": (["outer_diameter", "distal_outer_diameter", "proximal_outer_diameter",
            "middle_outer_diameter", "diameter"], "D8E4BC"),                # green
    "length": (["length", "stent_length", "balloon_length"], "FDE9D9"),     # orange
    "size": (["coil_diameter", "stent_diameter", "balloon_diameter", "implant_size",
              "wire_diameter", "vessel_diameter"], "E4DFEC"),               # purple
    "compat": (["guidewire_compatibility", "catheter_compatibility"], "FFF2CC"),  # yellow
    "provenance": (["evidence_layer", "source_year", "locator"], "E7E6E6"),  # grey
}
COLUMN_FILL = {column: fill for _, (columns, fill) in GROUPS.items() for column in columns}
SPEC_ORDER = [column for _, (columns, _) in GROUPS.items() for column in columns]

HEADER_LABELS = {
    "inner_diameter": "ID", "outer_diameter": "OD",
    "distal_outer_diameter": "Distal OD", "proximal_outer_diameter": "Proximal OD",
    "middle_outer_diameter": "Middle OD", "diameter": "Diameter",
    "length": "Length", "stent_length": "Stent Length", "balloon_length": "Balloon Length",
    "coil_diameter": "Coil Diameter", "stent_diameter": "Stent Diameter",
    "balloon_diameter": "Balloon Diameter", "implant_size": "Implant Size",
    "wire_diameter": "Wire Diameter", "vessel_diameter": "Vessel Diameter",
    "guidewire_compatibility": "Guidewire Compat.", "catheter_compatibility": "Catheter Compat.",
    "evidence_layer": "Evidence Layer", "source_year": "Year", "locator": "Locator",
    "device_name": "Device", "manufacturer": "Manufacturer",
    "catalog_number": "Catalog No.", "product_name": "Product / Variant",
    "labeled_statement": "Labeled Statement (as printed)",
}

WIDTHS = {"device_name": 30, "manufacturer": 14, "catalog_number": 18, "product_name": 30,
          "locator": 22, "labeled_statement": 90}


def sheet_title(category: str) -> str:
    title = re.sub(r"[\\/*?:\[\]]", "-", category)[:31]
    return title or "uncategorized"


def category_columns(rows: list[dict]) -> tuple[list[str], bool]:
    """Spec columns with at least one value in this category, plus prose flag."""
    used = [column for column in SPEC_ORDER
            if column in COLUMN_FILL and any(str(r.get(column) or "").strip() for r in rows)]
    has_prose = any(str(r.get("labeled_statement") or "").strip() for r in rows)
    return used, has_prose


def write_sheet(book: Workbook, category: str, rows: list[dict]) -> None:
    sheet = book.create_sheet(sheet_title(category))
    spec_columns, has_prose = category_columns(rows)
    columns = IDENTITY_COLUMNS + spec_columns + (["labeled_statement"] if has_prose else [])

    header_font = Font(bold=True)
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=HEADER_LABELS.get(column, column))
        cell.font = header_font
        fill = COLUMN_FILL.get(column)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = WIDTHS.get(column, 16)

    wrap = Alignment(wrap_text=True, vertical="top")
    ordered = sorted(rows, key=lambda r: (r.get("device_name", ""), r.get("catalog_number", "")))
    for row_index, row in enumerate(ordered, start=2):
        for column_index, column in enumerate(columns, start=1):
            value = str(row.get(column) or "")
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            fill = COLUMN_FILL.get(column)
            if fill and value:
                cell.fill = PatternFill("solid", fgColor=fill)
            if column == "labeled_statement":
                cell.alignment = wrap

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(ordered) + 1}"


def write_legend(book: Workbook) -> None:
    sheet = book.create_sheet("legend", 0)
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 70
    entries = [
        ("id", "Inner diameter (ID)"),
        ("od", "Outer diameters (OD, distal/proximal/middle, generic diameter)"),
        ("length", "Lengths (usable, stent, balloon)"),
        ("size", "Implant and device sizes (coil, stent, balloon, implant, wire, vessel)"),
        ("compat", "Compatibility (guidewire, catheter)"),
        ("provenance", "Provenance (evidence layer, source year, page locator)"),
    ]
    sheet["A1"], sheet["B1"] = "Color", "Meaning"
    sheet["A1"].font = sheet["B1"].font = Font(bold=True)
    for index, (group, meaning) in enumerate(entries, start=2):
        _, fill = GROUPS[group]
        cell = sheet.cell(row=index, column=1, value=group)
        cell.fill = PatternFill("solid", fgColor=fill)
        sheet.cell(row=index, column=2, value=meaning)
    sheet.cell(row=len(entries) + 3, column=1, value="Values are as printed in the cited source; every row carries its evidence layer, year, and page locator. One sheet per device category; columns a category never uses are omitted.")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()

    rows = build_rows()
    findings = validate_csv_rows(rows)
    summary = summarize(findings)
    if summary["errors"]:
        raise SystemExit(f"workbook blocked: {summary['errors']} validation errors")

    by_category: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        by_category[row.get("category") or "uncategorized"].append(row)

    book = Workbook()
    book.remove(book.active)
    write_legend(book)
    for category in sorted(by_category):
        write_sheet(book, category, by_category[category])

    args.out.parent.mkdir(parents=True, exist_ok=True)
    book.save(args.out)
    print(json.dumps({"sheets": len(by_category) + 1, "rows": len(rows),
                      "out": str(args.out)}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
